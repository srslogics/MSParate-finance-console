"""Full HTTP + PostgreSQL smoke suite. Run ONLY against a fresh disposable local DB.
Start uvicorn with DATABASE_URL explicitly pointing at that DB before running.
This script refuses non-loopback targets and refuses an already initialized app.
"""
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
import os
from urllib.parse import urlparse
from uuid import uuid4
import httpx
from openpyxl import load_workbook

BASE = os.environ.get('AUDIT_BASE_URL', 'http://127.0.0.1:8019')
assert urlparse(BASE).hostname in ('localhost', '127.0.0.1'), 'Disposable local app only'
client = httpx.Client(base_url=BASE, timeout=60)
assert client.get('/auth/setup-status').json() == {'has_users': False}, 'Use a fresh disposable database'
DAY = '2026-09-29'

def call(method, path, payload=None, headers=None, status=200, **kwargs):
    r = client.request(method, path, json=payload, headers=headers, **kwargs)
    assert r.status_code == status, (path, r.status_code, r.text[:300])
    if status == 200 and r.headers.get('content-type', '').startswith('application/json'):
        data = r.json()
        assert not (isinstance(data, dict) and data.get('error')), (path, data)
        return data
    return r

assert client.get('/dashboard').status_code == 401
assert client.get('/templates/vendor').status_code == 401
assert client.get('/.env').status_code == 401
assert client.get('/').status_code == 200
assert client.get('/healthz').json()['status'] == 'ok'
setup = call('POST', '/auth/setup-owner', {'username':'audit_owner','password':'local-test-password','display_name':'Synthetic audit owner'})
client.headers['X-Auth-Token'] = setup['token']
outlet = setup['user']['outlets'][0]['id']
client.headers['X-Outlet-Id'] = outlet
other = call('POST','/outlets',{'name':'Synthetic second outlet','code':'AUDIT2'})
other_id = other['outlet']['id']
call('POST','/auth/users',{'username':'audit_staff','password':'local-test-password','role':'STAFF','outlet_ids':[outlet]})
staff = call('POST','/auth/login',{'username':'audit_staff','password':'local-test-password'})
staff_headers={'X-Auth-Token':staff['token'],'X-Outlet-Id':outlet}
call('GET','/auth/users',headers=staff_headers,status=403)
call('GET',f'/dashboard?date={DAY}',headers={**staff_headers,'X-Outlet-Id':other_id},status=403)
print('PASS authentication and staff/outlet access controls')

def draft(**changes):
    return dict({'request_id':str(uuid4()),'date':DAY,'customer_name':'Synthetic audit customer','payment_mode':'Cash','paid_amount':4.25,'items':[{'item_name':'LIVER POT','line_type':'DRESSED','unit':'KGS','weight':0.1,'rate':100,'amount':10}]}, **changes)

payload=draft()
with ThreadPoolExecutor(max_workers=5) as pool:
    replies=list(pool.map(lambda _:call('POST','/retail-bills',payload),range(5)))
assert len({r['bill']['id'] for r in replies})==1
bill=replies[0]['bill'];bill_id=bill['id']
assert bill['total_amount']==10 and bill['outstanding_amount']==5.75 and bill['paid_amount']==4.25
assert bill['party_balance']==5.75,bill
assert bill['time'][:2] != '',bill
call('POST','/retail-bills',{**payload,'paid_amount':2},status=409)
assert call('POST','/retail-bills',payload)['bill']['id']==bill_id
call('PUT',f'/retail-bills/{bill_id}',{**payload,'bill_number':bill['bill_number']},headers=staff_headers,status=403)
assert client.get(f'/retail-bills/{bill_id}',headers={'X-Outlet-Id':other_id}).json().get('error')
# Updating a bill must replace old ledger effects, not append duplicates.
updated=call('PUT',f'/retail-bills/{bill_id}',{**payload,'bill_number':bill['bill_number'],'paid_amount':2})['bill']
assert updated['party_balance']==8,updated
receipt_payload={'request_id':str(uuid4()),'date':DAY,'party_name':'Synthetic audit customer','amount':3,'direction':'RECEIVED','payment_mode':'Cash'}
with ThreadPoolExecutor(max_workers=4) as pool:
    receipts=list(pool.map(lambda _:call('POST','/payment-receipts',receipt_payload),range(4)))
assert len({r['receipt']['id'] for r in receipts})==1
assert receipts[0]['receipt']['balance_after']==5,receipts[0]
assert call('GET',f'/party/profile?name=Synthetic%20audit%20customer')['party']['balance_after']==5
# Concurrent distinct transactions use the shared per-day document sequence.
with ThreadPoolExecutor(max_workers=6) as pool:
    bills=list(pool.map(lambda _:call('POST','/retail-bills',draft(customer_name='',paid_amount=10))['bill'],range(6)))
assert len({r['bill_number'] for r in bills} | {bill['bill_number'],receipts[0]['receipt']['receipt_number']})==8
print('PASS duplicate/concurrent saves, shared numbering, bill edits and payment ledger balances')
for value in ['NaN','Infinity','-Infinity','abc','1e9999']:
    call('POST','/retail-bills',draft(paid_amount=value),status=422)
for field in ['rate','weight','amount']:
    bad=draft();bad['items'][0][field]=-1
    assert client.post('/retail-bills',json=bad).json().get('error')
assert client.post('/retail-bills',json=draft(customer_name='',paid_amount=0)).json().get('error')
assert client.post('/retail-bills',json=draft(items=[{'item_name':'Zero','quantity':1,'rate':0}])).json().get('error')
paise=call('POST','/retail-bills',draft(customer_name='',paid_amount=None,items=[{'item_name':'Synthetic fraction','line_type':'DRESSED','weight':0.333,'rate':101}]))['bill']
assert paise['total_amount']==33.63 and paise['paid_amount']==33.63,paise
print('PASS invalid numbers, negative inputs, credit validation and paise precision')
# Preview must roll back. Same source file is valid independently at two outlets.
files={
 'vendor':'DATE,VENDOR,KGS,RATE,HEN_TYPE,NAG\n29/09/2026,Synthetic wholesale customer,5,100,BB,2\n',
 'dealer':'DATE,DEALER,KGS,RATE,HEN_TYPE,NAG\n29/09/2026,Synthetic supplier,20,80,BB,10\n',
 'payment':'DATE,PARTY,AMOUNT,PAYMENT_MODE,DIRECTION\n29/09/2026,Synthetic supplier,100,Cash,PAID\n',
 'opening-balance':'DATE,PARTY,OPENING_BALANCE,BALANCE_TYPE\n29/09/2026,Synthetic opening party,250,RECEIVABLE\n',
 'opening-stock':'DATE,HEN_TYPE,OPENING_KGS,OPENING_NAG\n28/09/2026,BB,100,50\n'
}
for kind,csv in files.items():
    path=f'/upload/{kind}'
    def upload(suffix='',headers=None):
        return client.post(path+suffix,files={'file':('synthetic.csv',csv,'text/csv')},headers=headers).json()
    preview=upload('?preview=true');assert preview.get('rows_inserted')==1,(kind,preview)
    saved=upload();assert saved.get('rows_inserted')==1,(kind,saved)
    assert upload().get('error')=='File already uploaded'
    second=upload(headers={'X-Outlet-Id':other_id});assert second.get('rows_inserted')==1,(kind,second)
    assert client.get(f'/templates/{kind}').status_code==200
profile=call('GET','/party/profile?name=Synthetic%20supplier')
assert abs(profile['party']['balance_after'])==1500,profile
second_profile=call('GET','/party/profile?name=Synthetic%20supplier',headers={'X-Outlet-Id':other_id})
assert abs(second_profile['party']['balance_after'])==1500,second_profile
print('PASS all 5 import types, preview rollback, duplicate detection and outlet isolation')
for path in [f'/dashboard?date={DAY}',f'/daily-sheet?date={DAY}&sheet_type=stock',f'/inventory/by-item?date={DAY}',f'/analytics/summary?start_date={DAY}&end_date={DAY}',f'/analytics/trend?start_date={DAY}&end_date={DAY}',f'/analytics/item-volume?start_date={DAY}&end_date={DAY}',f'/analytics/payment-modes?start_date={DAY}&end_date={DAY}',f'/top-debtors?start_date={DAY}&end_date={DAY}', '/top-payables',f'/retail-bills?date={DAY}',f'/payment-receipts?date={DAY}']:
    call('GET',path)
xlsx=call('GET',f'/daily-sheet/export?date={DAY}&sheet_type=stock')
wb=load_workbook(BytesIO(xlsx.content));assert wb.active.max_row>5
pdf=call('GET',f'/reports/export?report_type=summary&file_format=pdf&start_date={DAY}&end_date={DAY}')
assert pdf.content.startswith(b'%PDF'),pdf.text[:100]
call('POST','/auth/logout',headers=staff_headers)
call('GET',f'/dashboard?date={DAY}',headers=staff_headers,status=401)
print('PASS dashboard, stock sheets, analytics, histories, Excel export and logout')
for _ in range(10):
    assert client.post('/auth/login',json={'username':'audit_throttle','password':'incorrect'}).json().get('error')
assert client.post('/auth/login',json={'username':'audit_throttle','password':'incorrect'}).status_code==429
print('PASS failed-login throttling')
print('ALL DISPOSABLE POSTGRESQL HTTP CHECKS PASSED')
